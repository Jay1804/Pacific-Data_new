"""Tracker pipeline: Bridge DB queries + iBridge export from the MIS portal."""
import io, os, re, zipfile
import pandas as pd, pymysql, requests

ENV_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), '.env')
ENV = dict(l.strip().split('=', 1) for l in open(ENV_PATH) if '=' in l and not l.startswith('#') and l.strip())
DEFAULT_PROCESS = "Sapient Mandate"


def conn():
    return pymysql.connect(host=ENV['DB_HOST'], user=ENV['DB_USER'], password=ENV['DB_PASSWORD'],
                           database=ENV['DB_NAME'], connect_timeout=30, read_timeout=1800)


def run(c, sql, params=()):
    with c.cursor() as cur:
        cur.execute(sql, params)
        return cur.fetchall(), [d[0] for d in cur.description]


def chunks(l, n=400):
    for i in range(0, len(l), n):
        yield l[i:i + n]


def ph(n):
    return ','.join(['%s'] * n)


ARS_Q = """SELECT DISTINCT case_ars_no
FROM ec_case_master ecm
INNER JOIN ec_master_company emc ON ecm.client_id=emc.company_id
INNER JOIN ec_client ec ON emc.company_id=ec.client_id
LEFT JOIN ec_client_process ecp ON ecm.process_id=ecp.process_id
INNER JOIN ec_case_candidates ecc1 ON ecc1.candidate_id=ecm.candidate_id
WHERE ec.client_external_id IN ({}) AND first_name <>'dummy' AND process_name=%s"""

# Advance Tracker query, unchanged apart from the ARS filter being parameterised
TRACKER_Q = """SELECT ecc.Case_Check_id,case_ars_no,emc.Company_name,
CONCAT(first_name,' ',IFNULL(middle_name,''),' ',IFNULL(last_name,'')) AS Candidate_name,
Process_name, received_date AS case_received_date,
checkpoint_live.fn_case_status(case_status) AS 'Case_status',
checkpoint_live.fn_check_status(check_status) AS 'check_status',
ec.check_name AS 'Check_unique_name',
check_severity,
REPLACE(ecc.closure_comments,'rn','') AS closure_comments,ecc.check_closure_date,
ecc.insuff_remarks,
(SELECT action_taken_on FROM ec_case_history ech WHERE ecc.case_check_id=ech.check_id
AND action_taken in (' case Status changed to : Case Insufficient',' case Status changed to : InSufficient','Insuff Raised','Marked Insufficient','Marked Insufficient (Parallel Research)','New Status - Case Insufficient','New Status - InSufficient','Vendor request closed & Insuff raised','Case Insufficient','Check Created | Marked Insufficient','Check Updated | Marked Insufficient','New Status - Case Insuff Updated','New Status - New Case | Case Insufficient','Check Insuff raised','Case level Insuff raised','New Status - Case Insufficiency raised','Insufficient - Rework on Report','Insuff accepted' ) ORDER BY action_id LIMIT 1) AS 'First_Insuff_Date',
ecc.insuff_fulfill_date,
office_name location,
IF(ecc.family_id=4,emei.institute_name,IF(ecc.family_id=5,emc1.company_name,emcy.city_name)) AS verification_source,
ecm.case_expected_closure_date case_due_date,check_created_on,ecc.go_ahead_date,
IFNULL(ec.CHECK_OPS_NAME,LEFT(REPLACE(family_name,' Family',''),3)) AS 'Check Ops Name'
FROM ec_case_master ecm
LEFT JOIN ec_case_checks ecc ON ecm.case_id = ecc.case_id
LEFT JOIN ec_master_company emc ON emc.company_id = ecm.client_id
LEFT JOIN ec_client_process ecp ON ecm.process_id=ecp.process_id
LEFT JOIN ec_case_candidates ecc1 ON ecc1.candidate_id=ecm.candidate_id
LEFT JOIN ec_master_company_locations emcl ON ecm.client_office_id=emcl.office_id
LEFT JOIN ec_case_check_verification_source eccvs ON ecc.case_check_id=eccvs.case_check_id
LEFT JOIN ec_master_company emc1 ON eccvs.org_id=emc1.company_id
LEFT JOIN ec_master_educational_institute emei ON eccvs.org_id=emei.institute_id
LEFT JOIN ec_master_city emcy ON emcy.city_id=eccvs.org_id
LEFT JOIN ec_checks ec ON ecc.check_id=ec.check_id
LEFT JOIN ec_check_families ecf ON ec.family_id=ecf.family_id
WHERE case_status NOT IN (8,14) AND check_status <> 9 AND ecc.check_id<>193
AND ecm.case_ars_no IN ({})"""

REPORT_Q = """SELECT case_ars_no, report_sent_on, report_type, report_severity
FROM ec_case_reports ecr JOIN ec_case_master ecm ON ecr.case_id=ecm.case_id
WHERE ecm.case_ars_no IN ({}) AND report_status=5"""

CONTACT_Q = """SELECT case_ars_no, EMAIL_ADDRESS,
IF(LENGTH(mobile_no)=0,contact_value,CONCAT('con ',MOBILE_NO)) AS MOBILE_NO
FROM ec_case_master ecm
LEFT JOIN ec_case_candidates ecc1 ON ecc1.candidate_id=ecm.candidate_id
LEFT JOIN ec_candidate_contact_details eccd ON ecm.candidate_id=eccd.candidate_id
WHERE case_ars_no LIKE '%%-%%' AND case_ars_no IN ({})"""

TYPE_NAME = {0: 'Interim', 1: 'Final', 2: 'Additional'}


def fetch_all(c, sql, ars):
    frames = []
    for ch in chunks(ars):
        rows, cols = run(c, sql.format(ph(len(ch))), ch)
        frames.append(pd.DataFrame(rows, columns=cols))
        print(f'  {len(frames[-1])} rows', flush=True)
    return pd.concat(frames, ignore_index=True)


def fetch_ibridge(ars, user, password, chunk=1500, progress=print):
    """Case initiation data - ARS Wise from the MIS portal (iBridge DB, no direct access)."""
    base = 'https://mis.authbridge.com/export_query/'
    s = requests.Session()
    r = s.post(base + 'login.php', data={'username': user, 'password': password, 'login': 'Login'})
    if 'home.php' not in r.url:
        raise RuntimeError('MIS portal login failed - check the MIS user ID / password')
    frames = []
    for ch in chunks(list(ars), chunk):
        data = {'hostname': '14', 'database': 'ibridge_ars', 'access_time': 'iBridge', 'csv_query': '1856',
                'query_days_range': '0', 'check_id1': ','.join("'%s'" % a for a in ch)}
        r = s.post(base + 'process.php', data=data, files={'x': (None, '')}, timeout=600)
        z = zipfile.ZipFile(io.BytesIO(r.content))
        frames.append(pd.read_csv(z.open(z.namelist()[0])))
        progress(f'iBridge: {len(frames[-1])} rows')
    return pd.concat(frames, ignore_index=True)


def build(client_ids, process_name=DEFAULT_PROCESS, mis_user=None, mis_password=None, ib_df=None, progress=print):
    """Returns (tracker DataFrame, list of ARS numbers)."""
    client_ids = [int(x) for x in client_ids]
    c = conn()
    rows, _ = run(c, ARS_Q.format(ph(len(client_ids))), (*client_ids, process_name))
    ars = sorted({r[0] for r in rows if r[0]})
    progress(f'{len(ars)} ARS numbers found')
    if not ars:
        return pd.DataFrame(), ars

    progress('Advance tracker'); t = fetch_all(c, TRACKER_Q, ars)
    progress('Reports'); r = fetch_all(c, REPORT_Q, ars)
    progress('Candidate contacts'); ct = fetch_all(c, CONTACT_Q, ars)
    r['report_sent_on'] = pd.to_datetime(r['report_sent_on'], errors='coerce')

    # Sent Cases1: latest sent report of any type
    latest = r.sort_values('report_sent_on', ascending=False).drop_duplicates('case_ars_no')
    latest = latest.assign(**{'Report Type': latest['report_type'].map(TYPE_NAME)})[
        ['case_ars_no', 'report_sent_on', 'Report Type', 'report_severity']]
    # Sent Cases2: oldest sent report
    first = r.sort_values('report_sent_on').drop_duplicates('case_ars_no')[['case_ars_no', 'report_sent_on']]         .rename(columns={'report_sent_on': 'First Report Sent Date'})

    def last_of(tp, name):
        x = r[r.report_type == tp].sort_values('report_sent_on', ascending=False).drop_duplicates('case_ars_no')
        return x[['case_ars_no', 'report_sent_on']].rename(columns={'report_sent_on': name})

    ct = ct.drop_duplicates('case_ars_no').rename(columns={'MOBILE_NO': 'Candidate Contact No.'})

    out = t.merge(latest, on='case_ars_no', how='left')
    out = out.merge(ct[['case_ars_no', 'EMAIL_ADDRESS', 'Candidate Contact No.']], on='case_ars_no', how='left')
    out = out.merge(first, on='case_ars_no', how='left')
    out = out.merge(last_of(1, 'L_Final Report Date'), on='case_ars_no', how='left')
    out = out.merge(last_of(2, 'L_Additional Report Date'), on='case_ars_no', how='left')

    if ib_df is None and mis_user and mis_password:
        progress('Fetching iBridge case initiation dates from MIS portal')
        ib_df = fetch_ibridge(ars, mis_user, mis_password, progress=progress)
    if ib_df is not None:
        norm = lambda x: re.sub(r'[^a-z]', '', str(x).lower())
        key = next(col for col in ib_df.columns if norm(col) in ('arsno', 'casearsno', 'arsnumber'))
        dt = next(col for col in ib_df.columns if norm(col) == 'caseinitiationdate')
        ib = ib_df[[key, dt]].drop_duplicates(key).rename(
            columns={key: 'case_ars_no', dt: 'iBridge Case Initiation Date'})
        out = out.merge(ib, on='case_ars_no', how='left')
    else:
        out['iBridge Case Initiation Date'] = None
    out['case_ars_no'] = out['case_ars_no'].astype(str)
    for c in DATE_COLS:
        out[c] = pd.to_datetime(out[c], errors='coerce').dt.normalize()
    return out, ars


DATE_COLS = ['report_sent_on', 'First Report Sent Date', 'L_Final Report Date',
             'L_Additional Report Date', 'iBridge Case Initiation Date']
DATE_FMT = 'dd-mmm-yyyy'


def to_excel_bytes(df):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    df = df.copy()
    for c in DATE_COLS:
        if c in df.columns:
            df[c] = pd.to_datetime(df[c], errors='coerce').dt.normalize()

    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine='openpyxl') as xw:
        df.to_excel(xw, index=False)
        ws = xw.sheets['Sheet1']
        cols = list(df.columns)
        side = Side(style='thin', color='000000')
        border = Border(left=side, right=side, top=side, bottom=side)
        head_fill = PatternFill('solid', start_color='D9D9D9', end_color='D9D9D9')
        date_idx = {cols.index(c) + 1 for c in DATE_COLS if c in cols}
        ars_idx = cols.index('case_ars_no') + 1

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = head_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                if cell.column in date_idx:
                    cell.number_format = DATE_FMT
                    cell.alignment = Alignment(horizontal='center')
                elif cell.column == ars_idx:
                    cell.number_format = '@'  # keep ARS as text

        for i, name in enumerate(cols, 1):
            longest = max([len(str(name))] + [len(str(v)) for v in df[name].head(500).dropna()])
            ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 12), 45)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    return bio.getvalue()
